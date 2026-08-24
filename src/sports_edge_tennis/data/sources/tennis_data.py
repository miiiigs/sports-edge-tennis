from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

import openpyxl

from ..contracts import OddsPriceRecord, SourceDocument
from ..hashing import source_hash
from ..normalize import collapse_whitespace


class TennisDataRowError(ValueError):
    """Raised when a Tennis-Data row cannot be safely normalized."""


_REQUIRED_COLUMNS = {"Date", "Tournament", "Winner", "Loser"}
_BOOKMAKER_EXCLUSIONS = {
    "W",
    "L",
    "WRank",
    "LRank",
    "WPts",
    "LPts",
    "Bof",
}
_SURFACE_MAP = {
    "Hard": "HARD",
    "Clay": "CLAY",
    "Grass": "GRASS",
    "Carpet": "CARPET",
}


@dataclass(frozen=True)
class TennisDataParseOutcome:
    records: list[OddsPriceRecord]
    rejected_rows: list[dict[str, object]]


def _iter_rows(path: Path) -> tuple[list[str], Iterable[list[object]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            rows = list(reader)
        return [str(value) for value in rows[0]], rows[1:]
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [str(value) if value is not None else "" for value in next(iterator)]
        return header, (list(row) for row in iterator)
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(path.as_posix())
        sheet = workbook.sheet_by_index(0)
        header = [str(sheet.cell_value(0, column)) for column in range(sheet.ncols)]
        date_column_indexes = {
            index for index, column_name in enumerate(header) if column_name.strip() == "Date"
        }

        def iter_xls_rows() -> Iterable[list[object]]:
            for row_index in range(1, sheet.nrows):
                row_values: list[object] = []
                for column in range(sheet.ncols):
                    cell = sheet.cell(row_index, column)
                    value = cell.value
                    if column in date_column_indexes and cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode).date()
                    row_values.append(value)
                yield row_values

        return header, iter_xls_rows()
    raise TennisDataRowError("unsupported_tennis_data_file_type")


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise TennisDataRowError("invalid_odds_date")


def _parse_decimal(value: object) -> Decimal | None:
    if value in ("", None):
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation as exc:
        raise TennisDataRowError("invalid_decimal_price") from exc


def _bookmaker_prefixes(headers: list[str]) -> list[str]:
    header_set = set(headers)
    prefixes: list[str] = []
    for header in headers:
        if not header.endswith("W"):
            continue
        prefix = header[:-1]
        if not prefix or prefix in _BOOKMAKER_EXCLUSIONS:
            continue
        if prefix.isdigit():
            continue
        loser_header = f"{prefix}L"
        if loser_header in header_set:
            prefixes.append(prefix)
    return sorted(set(prefixes))


def _surface(value: object) -> str | None:
    text = str(value).strip()
    if not text:
        return None
    normalized = _SURFACE_MAP.get(text)
    if normalized is None:
        raise TennisDataRowError("unsupported_surface")
    return normalized


def _source_observed_at(event_date: date) -> datetime:
    return datetime(event_date.year, event_date.month, event_date.day, 23, 59, 59, tzinfo=UTC)


def parse_tennis_data_file(document: SourceDocument) -> TennisDataParseOutcome:
    records: list[OddsPriceRecord] = []
    rejected_rows: list[dict[str, object]] = []
    headers, rows = _iter_rows(document.local_path)
    missing = sorted(_REQUIRED_COLUMNS - set(headers))
    if missing:
        raise TennisDataRowError(f"missing_required_columns:{','.join(missing)}")

    prefixes = _bookmaker_prefixes(headers)
    for row_number, raw_row in enumerate(rows, start=2):
        row = {headers[index]: raw_row[index] for index in range(len(headers))}
        try:
            event_date = _parse_date(row["Date"])
            winner_name = collapse_whitespace(str(row["Winner"]))
            loser_name = collapse_whitespace(str(row["Loser"]))
            tournament_name = collapse_whitespace(str(row["Tournament"]))
            if not winner_name or not loser_name:
                raise TennisDataRowError("missing_player_name")
            if not tournament_name:
                raise TennisDataRowError("missing_tournament_name")

            source_observed_at = _source_observed_at(event_date)
            for prefix in prefixes:
                winner_price = _parse_decimal(row.get(f"{prefix}W"))
                loser_price = _parse_decimal(row.get(f"{prefix}L"))
                if winner_price is None and loser_price is None:
                    continue
                if winner_price is None or loser_price is None:
                    raise TennisDataRowError("half_populated_bookmaker_pair")
                for selection_name, selection_role, price in (
                    (winner_name, "WINNER", winner_price),
                    (loser_name, "LOSER", loser_price),
                ):
                    if price <= Decimal("1"):
                        raise TennisDataRowError("invalid_decimal_price")
                    raw_fields = {
                        "source_row_number": row_number,
                        "bookmaker_code": prefix,
                        "selection_role": selection_role,
                        "selection_name": selection_name,
                        "winner_name": winner_name,
                        "loser_name": loser_name,
                        "event_date": event_date.isoformat(),
                        "tournament_name": tournament_name,
                        "price": price,
                    }
                    records.append(
                        OddsPriceRecord(
                            document=document,
                            source_row_number=row_number,
                            source_snapshot_id=f"{document.filename}#{row_number}:{prefix}",
                            source_record_hash=source_hash(document.source_system, raw_fields),
                            source_record_locator=f"{document.filename}#{row_number}:{prefix}",
                            tournament_name=tournament_name,
                            event_date=event_date,
                            surface=_surface(row.get("Surface")),
                            round_code=collapse_whitespace(str(row["Round"])) if row.get("Round") else None,
                            best_of_sets=int(row["Best of"]) if row.get("Best of") not in ("", None) else None,
                            winner_name=winner_name,
                            loser_name=loser_name,
                            selection_name=selection_name,
                            selection_role=selection_role,
                            bookmaker_code=prefix,
                            decimal_price=price,
                            source_observed_at=source_observed_at,
                            as_of=source_observed_at,
                            raw_fields={key: value for key, value in row.items() if value not in ("", None)},
                            limitations=(
                                "Tennis-Data historical odds are treated as date-level closing-style observations because the archive does not expose reliable intraday timestamps here.",
                            ),
                        )
                    )
        except (TennisDataRowError, ValueError, TypeError) as exc:
            rejected_rows.append(
                {
                    "row_number": row_number,
                    "reason": str(exc),
                    "tournament_name": row.get("Tournament"),
                    "winner_name": row.get("Winner"),
                    "loser_name": row.get("Loser"),
                }
            )
    return TennisDataParseOutcome(records=records, rejected_rows=rejected_rows)
